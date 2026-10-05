import csv
import io
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import lab_finder as lf  # noqa: E402


class FakeResponse:
    def __init__(self, status_code, payload):
        self.status_code = status_code
        self._payload = payload
        self.text = str(payload)

    def json(self):
        return self._payload


class FakeSession:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def post(self, url, json, headers, timeout):
        self.calls.append({"url": url, "json": json, "headers": headers})
        return self.responses.pop(0)


def make_place(pid, name="Ege Diş Protez Laboratuvarı", district="Karşıyaka", province="İzmir",
               lat=38.46, lng=27.11, status="OPERATIONAL", **extra):
    place = {
        "id": pid,
        "displayName": {"text": name, "languageCode": "tr"},
        "formattedAddress": f"1234. Sk. No:5, 35590 {district}/{province}, Türkiye",
        "addressComponents": [
            {"longText": district, "shortText": district, "types": ["administrative_area_level_2", "political"]},
            {"longText": province, "shortText": province, "types": ["administrative_area_level_1", "political"]},
        ],
        "location": {"latitude": lat, "longitude": lng},
        "businessStatus": status,
    }
    place.update(extra)
    return place


class NormalizeTest(unittest.TestCase):
    def test_turkish_letters(self):
        self.assertEqual(lf.normalize("Karşıyaka"), "karsiyaka")
        self.assertEqual(lf.normalize("İZMİR"), "izmir")
        self.assertEqual(lf.normalize("Çiğli Ödemiş Güzelbahçe"), "cigli odemis guzelbahce")
        self.assertEqual(lf.normalize(None), "")


class ClassifyTest(unittest.TestCase):
    def test_dental_labs(self):
        for name in ["Ege Diş Protez Laboratuvarı", "Dentlab İzmir", "Zirkonyum Diş Teknisyeni",
                     "Porselen Dental Lab", "Kuzey CAD/CAM Merkezi", "Tıbbi Diş Protez Lab"]:
            self.assertTrue(lf.is_likely_lab(name), name)

    def test_not_labs(self):
        for name in ["Karşıyaka Ağız ve Diş Sağlığı Polikliniği", "Dr. Ayşe Yılmaz Diş Hekimi",
                     "Özel Tıbbi Tahlil Laboratuvarı", "Ege Gıda Analiz Laboratuvarı"]:
            self.assertFalse(lf.is_likely_lab(name), name)


class AddressTest(unittest.TestCase):
    def test_district_from_components(self):
        self.assertEqual(lf.district_of(make_place("a", district="Bornova"), "Konak"), "Bornova")

    def test_district_from_address_fallback(self):
        place = {"formattedAddress": "Bostanlı, 1234. Sk. No:5, 35590 Karşıyaka/İzmir, Türkiye"}
        self.assertEqual(lf.district_of(place, "Konak"), "Karşıyaka")
        self.assertEqual(lf.district_of({}, "Konak"), "Konak")

    def test_is_in_izmir(self):
        self.assertTrue(lf.is_in_izmir(make_place("a")))
        self.assertFalse(lf.is_in_izmir(make_place("a", province="Manisa")))
        self.assertTrue(lf.is_in_izmir({"formattedAddress": "Konak/IZMIR, Türkiye"}))
        self.assertFalse(lf.is_in_izmir({"formattedAddress": "Yunusemre/Manisa, Türkiye"}))


class ClientTest(unittest.TestCase):
    def test_pagination_sends_token_and_headers(self):
        session = FakeSession([
            FakeResponse(200, {"places": [make_place("a")], "nextPageToken": "tok1"}),
            FakeResponse(200, {"places": [make_place("b")]}),
        ])
        client = lf.PlacesClient("KEY", session=session, sleep=lambda s: None)
        places = client.search_text("diş protez laboratuvarı Karşıyaka İzmir")
        self.assertEqual([p["id"] for p in places], ["a", "b"])
        self.assertEqual(client.request_count, 2)
        first, second = session.calls
        self.assertEqual(first["headers"]["X-Goog-Api-Key"], "KEY")
        self.assertIn("nextPageToken", first["headers"]["X-Goog-FieldMask"])
        self.assertNotIn("pageToken", first["json"])
        self.assertEqual(second["json"]["pageToken"], "tok1")
        self.assertEqual(second["json"]["textQuery"], first["json"]["textQuery"])
        self.assertEqual(first["json"]["pageSize"], 20)

    def test_stops_after_max_pages(self):
        session = FakeSession([FakeResponse(200, {"places": [make_place(str(i))], "nextPageToken": "t"})
                               for i in range(5)])
        client = lf.PlacesClient("KEY", session=session, sleep=lambda s: None)
        self.assertEqual(len(client.search_text("q")), lf.MAX_PAGES)

    def test_location_restriction(self):
        session = FakeSession([FakeResponse(200, {})])
        client = lf.PlacesClient("KEY", session=session, sleep=lambda s: None)
        rect = lf._rect(38.0, 27.0, 38.5, 27.5)
        self.assertEqual(client.search_text("q", rect), [])
        self.assertEqual(session.calls[0]["json"]["locationRestriction"], {"rectangle": rect})

    def test_retries_on_429(self):
        sleeps = []
        session = FakeSession([
            FakeResponse(429, {"error": {"message": "quota"}}),
            FakeResponse(200, {"places": [make_place("a")]}),
        ])
        client = lf.PlacesClient("KEY", session=session, sleep=sleeps.append)
        self.assertEqual(len(client.search_text("q")), 1)
        self.assertEqual(sleeps, [2])

    def test_gives_up_after_retries(self):
        session = FakeSession([FakeResponse(503, {"error": {"message": "down"}})] * (lf.MAX_RETRIES + 1))
        client = lf.PlacesClient("KEY", session=session, sleep=lambda s: None)
        with self.assertRaisesRegex(lf.PlacesError, "down"):
            client.search_text("q")

    def test_no_retry_on_403(self):
        session = FakeSession([FakeResponse(403, {"error": {"message": "API key not valid"}})])
        client = lf.PlacesClient("KEY", session=session, sleep=lambda s: None)
        with self.assertRaisesRegex(lf.PlacesError, "403.*API key not valid"):
            client.search_text("q")
        self.assertEqual(len(session.calls), 1)


class FakeClient:
    """search_area testləri üçün: ərazisiz sorğuya 60, bölünmüş ərazilərə 2 nəticə qaytarır."""

    def __init__(self, capped_depths=(None,)):
        self.calls = []
        self.capped_depths = capped_depths

    def search_text(self, text_query, rect=None):
        self.calls.append(rect)
        n = len(self.calls)
        if rect is None or "always" in self.capped_depths:
            return [make_place(f"{n}-{i}", lat=38.40 + i * 0.001, lng=27.10 + i * 0.001) for i in range(60)]
        return [make_place(f"{n}-x"), make_place(f"{n}-y")]


class SearchAreaTest(unittest.TestCase):
    def test_splits_when_capped(self):
        client = FakeClient()
        places = lf.search_area(client, "q")
        self.assertEqual(len(client.calls), 5)
        self.assertIsNone(client.calls[0])
        self.assertEqual(len(places), 60 + 4 * 2)
        lows = {(r["low"]["latitude"], r["low"]["longitude"]) for r in client.calls[1:]}
        self.assertEqual(len(lows), 4)

    def test_depth_limit(self):
        client = FakeClient(capped_depths=("always",))
        with redirect_stderr(io.StringIO()):
            lf.search_area(client, "q")
        self.assertEqual(len(client.calls), 1 + 4 + 16)

    def test_bounding_box_clamped_to_izmir(self):
        box = lf.bounding_box([make_place("a", lat=37.0, lng=26.0), make_place("b", lat=40.0, lng=29.0)])
        self.assertEqual(box, lf._rect(*lf.IZMIR_BOUNDS["low"], *lf.IZMIR_BOUNDS["high"]))
        self.assertIsNone(lf.bounding_box([{"id": "x"}]))

    def test_split_rect(self):
        quads = lf.split_rect(lf._rect(38.0, 27.0, 39.0, 28.0))
        self.assertEqual(quads[0], lf._rect(38.0, 27.0, 38.5, 27.5))
        self.assertEqual(quads[3], lf._rect(38.5, 27.5, 39.0, 28.0))


class CollectAndRowsTest(unittest.TestCase):
    def setUp(self):
        responses = {
            "diş protez laboratuvarı Karşıyaka İzmir": [
                make_place("lab1", nationalPhoneNumber="0232 111 11 11", websiteUri="https://ege.example",
                           googleMapsUri="https://maps.google.com/?cid=1", rating=4.8, userRatingCount=12),
                make_place("clinic", name="Karşıyaka Ağız ve Diş Sağlığı Polikliniği",
                           primaryTypeDisplayName={"text": "Diş Hekimi"}),
                make_place("closed", name="Eski Protez Lab", status="CLOSED_PERMANENTLY"),
                make_place("manisa", name="Manisa Dental Lab", district="Yunusemre", province="Manisa"),
            ],
            "dental laboratuvar Karşıyaka İzmir": [
                make_place("lab1"),
                make_place("lab2", name="Bostanlı Dental Lab", district="Çiğli", status="CLOSED_TEMPORARILY"),
            ],
        }

        class Client:
            def search_text(self, text_query, rect=None):
                return responses.get(text_query, [])

        queries = lf.build_queries(["Karşıyaka"], ["diş protez laboratuvarı", "dental laboratuvar"])
        with redirect_stderr(io.StringIO()):
            self.records = lf.collect(Client(), queries, {})

    def test_dedup_and_keywords(self):
        self.assertEqual(set(self.records), {"lab1", "clinic", "closed", "manisa", "lab2"})
        self.assertEqual(list(self.records["lab1"]["keywords"]),
                         ["diş protez laboratuvarı", "dental laboratuvar"])

    def test_build_rows(self):
        labs, review, skipped = lf.build_rows(self.records)
        self.assertEqual([r["name"] for r in labs], ["Bostanlı Dental Lab", "Ege Diş Protez Laboratuvarı"])
        self.assertEqual([r["name"] for r in review], ["Karşıyaka Ağız ve Diş Sağlığı Polikliniği"])
        self.assertEqual(skipped, {"closed": 1, "outside": 1})
        ege = labs[1]
        self.assertEqual(ege["phone"], "0232 111 11 11")
        self.assertEqual(ege["district"], "Karşıyaka")
        self.assertEqual(ege["keywords"], "diş protez laboratuvarı, dental laboratuvar")
        self.assertEqual(labs[0]["district"], "Çiğli")
        self.assertEqual(labs[0]["note"], "Müvəqqəti bağlı")
        self.assertEqual(review[0]["category"], "Diş Hekimi")

    def test_write_files(self):
        labs, review, _ = lf.build_rows(self.records)
        with tempfile.TemporaryDirectory() as tmp:
            xlsx, csv_path = Path(tmp, "out.xlsx"), Path(tmp, "out.csv")
            lf.write_xlsx(xlsx, labs, review)
            lf.write_csv(csv_path, labs, review)

            wb = load_workbook(xlsx)
            self.assertEqual(wb.sheetnames, ["Lablar", "Yoxlanmalı"])
            ws = wb["Lablar"]
            self.assertEqual(ws.cell(row=1, column=1).value, "Ad")
            self.assertEqual(ws.max_row, 3)
            self.assertEqual(ws.cell(row=3, column=8).hyperlink.target, "https://maps.google.com/?cid=1")
            self.assertEqual(wb["Yoxlanmalı"].max_row, 2)

            with open(csv_path, encoding="utf-8-sig", newline="") as f:
                rows = list(csv.reader(f))
            self.assertEqual(rows[0][:3], ["Növ", "Ad", "İlçə"])
            self.assertEqual([r[0] for r in rows[1:]], ["Lab", "Lab", "Yoxlanmalı"])


class CliTest(unittest.TestCase):
    def test_resolve_districts(self):
        self.assertEqual(lf.resolve_districts(["karsiyaka", "ÇİĞLİ"]), ["Karşıyaka", "Çiğli"])
        self.assertEqual(len(lf.resolve_districts(None)), 30)
        with self.assertRaises(SystemExit):
            lf.resolve_districts(["Kadıköy"])

    def test_dry_run(self):
        out = io.StringIO()
        with redirect_stdout(out):
            self.assertEqual(lf.main(["--dry-run", "--districts", "Karşıyaka"]), 0)
        lines = out.getvalue().splitlines()
        self.assertEqual(lines[0], "diş protez laboratuvarı Karşıyaka İzmir")
        self.assertIn("5 sorğu", out.getvalue())

    def test_missing_key(self):
        env = os.environ.pop("GOOGLE_MAPS_API_KEY", None)
        try:
            with redirect_stderr(io.StringIO()):
                self.assertEqual(lf.main(["--districts", "Karşıyaka"]), 1)
        finally:
            if env is not None:
                os.environ["GOOGLE_MAPS_API_KEY"] = env


if __name__ == "__main__":
    unittest.main()

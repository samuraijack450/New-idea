#!/usr/bin/env python3
"""İzmir diş protez laboratoriyalarını Google Places API (New) ilə tapır.

İstifadə:
    python lab_finder.py                                  # bütün İzmir
    python lab_finder.py --districts Karşıyaka Bornova    # seçilmiş ilçələr
    python lab_finder.py --dry-run                        # API-yə getmədən sorğuları göstər

API açarı GOOGLE_MAPS_API_KEY environment variable-dan oxunur.
"""

import argparse
import csv
import os
import re
import sys
import time
import unicodedata
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font

SEARCH_URL = "https://places.googleapis.com/v1/places:searchText"
FIELD_MASK = ",".join([
    "places.id",
    "places.displayName",
    "places.formattedAddress",
    "places.addressComponents",
    "places.location",
    "places.nationalPhoneNumber",
    "places.internationalPhoneNumber",
    "places.websiteUri",
    "places.rating",
    "places.userRatingCount",
    "places.googleMapsUri",
    "places.businessStatus",
    "places.primaryType",
    "places.primaryTypeDisplayName",
    "nextPageToken",
])
PAGE_SIZE = 20
MAX_PAGES = 3  # Google bir sorğu üçün ən çox 60 nəticə qaytarır
MAX_SPLIT_DEPTH = 2  # limitə çatan sorğu ərazisi neçə dəfə 4 hissəyə bölünə bilər
MAX_RETRIES = 4
RETRY_STATUSES = {429, 500, 502, 503, 504}

# İzmir vilayətinin təxmini sərhədləri (bölünən ərazilər bundan kənara çıxmır)
IZMIR_BOUNDS = {"low": (37.70, 26.10), "high": (39.50, 28.55)}

IZMIR_DISTRICTS = [
    "Aliağa", "Balçova", "Bayındır", "Bayraklı", "Bergama", "Beydağ",
    "Bornova", "Buca", "Çeşme", "Çiğli", "Dikili", "Foça",
    "Gaziemir", "Güzelbahçe", "Karabağlar", "Karaburun", "Karşıyaka", "Kemalpaşa",
    "Kınık", "Kiraz", "Konak", "Menderes", "Menemen", "Narlıdere",
    "Ödemiş", "Seferihisar", "Selçuk", "Tire", "Torbalı", "Urla",
]

KEYWORDS = [
    "diş protez laboratuvarı",
    "dental laboratuvar",
    "diş laboratuvarı",
    "protez laboratuvarı",
    "diş teknisyeni",
]

# Adına görə "lab" sayılan yerlər; qalanları "Yoxlanmalı" vərəqinə düşür
LAB_RE = re.compile(r"lab|protez|teknisyen|zirkon|seramik|porselen|cad ?/?cam")
DENTAL_RE = re.compile(r"\bdis\b|dental|protez|zirkon|seramik|porselen|ortodon|implant")
OTHER_LAB_RE = re.compile(
    r"tahlil|tibbi|biyokimya|patoloji|mikrobiyoloji|goruntuleme|rontgen|"
    r"veteriner|gida|analiz|kalibrasyon|cevre|hormon|genetik"
)

COLUMNS = [
    ("name", "Ad"),
    ("district", "İlçə"),
    ("address", "Ünvan"),
    ("phone", "Telefon"),
    ("website", "Sayt"),
    ("rating", "Reytinq"),
    ("reviews", "Rəy sayı"),
    ("maps_url", "Maps linki"),
    ("category", "Google kateqoriyası"),
    ("note", "Qeyd"),
    ("lat", "Lat"),
    ("lng", "Lng"),
    ("keywords", "Tapıldığı açar sözlər"),
]
LINK_FIELDS = {"website", "maps_url"}

_TR_MAP = str.maketrans("İIıŞşĞğÜüÖöÇçÂâÎîÛû", "iiissgguuooccaaiiuu")


class PlacesError(RuntimeError):
    pass


def log(message):
    print(message, file=sys.stderr, flush=True)


def normalize(text):
    """Türk hərflərini ASCII-yə çevirib kiçik hərfə salır: 'Karşıyaka' -> 'karsiyaka'."""
    text = (text or "").translate(_TR_MAP).lower()
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


class PlacesClient:
    def __init__(self, api_key, session=None, sleep=time.sleep):
        self.api_key = api_key
        self.session = session or requests.Session()
        self.sleep = sleep
        self.request_count = 0

    def _post(self, body):
        headers = {
            "Content-Type": "application/json",
            "X-Goog-Api-Key": self.api_key,
            "X-Goog-FieldMask": FIELD_MASK,
        }
        error = ""
        for attempt in range(MAX_RETRIES + 1):
            self.request_count += 1
            try:
                resp = self.session.post(SEARCH_URL, json=body, headers=headers, timeout=30)
            except requests.RequestException as exc:
                error = str(exc)
            else:
                if resp.status_code == 200:
                    return resp.json()
                error = _error_message(resp)
                if resp.status_code not in RETRY_STATUSES:
                    raise PlacesError(f"Google Places API səhvi ({resp.status_code}): {error}")
            if attempt < MAX_RETRIES:
                self.sleep(2 ** (attempt + 1))
        raise PlacesError(f"Google Places API cavab vermədi: {error}")

    def search_text(self, text_query, rect=None):
        """Bir mətn sorğusunun bütün səhifələrini (ən çox 60 nəticə) qaytarır."""
        body = {
            "textQuery": text_query,
            "languageCode": "tr",
            "regionCode": "TR",
            "pageSize": PAGE_SIZE,
        }
        if rect:
            body["locationRestriction"] = {"rectangle": rect}
        places = []
        for _ in range(MAX_PAGES):
            data = self._post(body)
            places.extend(data.get("places", []))
            token = data.get("nextPageToken")
            if not token:
                break
            body = {**body, "pageToken": token}
        return places


def _error_message(resp):
    try:
        return resp.json()["error"]["message"]
    except (ValueError, KeyError, TypeError):
        return resp.text[:300]


def _rect(low_lat, low_lng, high_lat, high_lng):
    return {
        "low": {"latitude": low_lat, "longitude": low_lng},
        "high": {"latitude": high_lat, "longitude": high_lng},
    }


def bounding_box(places, margin=0.25, min_pad=0.01):
    """Nəticələri əhatə edən, hər tərəfdən genişləndirilmiş və İzmir sərhədinə kəsilmiş düzbucaqlı."""
    coords = [
        (p["location"]["latitude"], p["location"]["longitude"])
        for p in places
        if "latitude" in p.get("location", {}) and "longitude" in p.get("location", {})
    ]
    if not coords:
        return None
    lats, lngs = zip(*coords)
    pad_lat = max((max(lats) - min(lats)) * margin, min_pad)
    pad_lng = max((max(lngs) - min(lngs)) * margin, min_pad)
    (bound_low_lat, bound_low_lng), (bound_high_lat, bound_high_lng) = (
        IZMIR_BOUNDS["low"], IZMIR_BOUNDS["high"],
    )
    return _rect(
        max(min(lats) - pad_lat, bound_low_lat),
        max(min(lngs) - pad_lng, bound_low_lng),
        min(max(lats) + pad_lat, bound_high_lat),
        min(max(lngs) + pad_lng, bound_high_lng),
    )


def split_rect(rect):
    low_lat, low_lng = rect["low"]["latitude"], rect["low"]["longitude"]
    high_lat, high_lng = rect["high"]["latitude"], rect["high"]["longitude"]
    mid_lat, mid_lng = (low_lat + high_lat) / 2, (low_lng + high_lng) / 2
    return [
        _rect(low_lat, low_lng, mid_lat, mid_lng),
        _rect(low_lat, mid_lng, mid_lat, high_lng),
        _rect(mid_lat, low_lng, high_lat, mid_lng),
        _rect(mid_lat, mid_lng, high_lat, high_lng),
    ]


def search_area(client, text_query, rect=None, depth=0):
    """Sorğu 60 nəticə limitinə çatırsa, ərazini 4 hissəyə bölüb hər birində yenidən axtarır."""
    places = client.search_text(text_query, rect)
    if len(places) < PAGE_SIZE * MAX_PAGES:
        return places
    if depth >= MAX_SPLIT_DEPTH:
        log(f"  xəbərdarlıq: '{text_query}' bölündükdən sonra da limitə çatır, bəzi nəticələr itə bilər")
        return places
    area = rect or bounding_box([p for p in places if is_in_izmir(p)])
    if area is None:
        return places
    for quadrant in split_rect(area):
        places.extend(search_area(client, text_query, quadrant, depth + 1))
    return places


def address_component(place, component_type):
    for comp in place.get("addressComponents", []):
        if component_type in comp.get("types", []):
            return comp.get("longText", "")
    return ""


def is_in_izmir(place):
    province = address_component(place, "administrative_area_level_1")
    if province:
        return normalize(province) == "izmir"
    return "izmir" in normalize(place.get("formattedAddress", ""))


def district_of(place, fallback):
    district = address_component(place, "administrative_area_level_2")
    if district:
        return district
    match = re.search(r"([^\s,/]+)/[İI]zmir", place.get("formattedAddress", ""))
    return match.group(1) if match else fallback


def is_likely_lab(name):
    n = normalize(name)
    if not LAB_RE.search(n):
        return False
    return not (OTHER_LAB_RE.search(n) and not DENTAL_RE.search(n))


def build_queries(districts, keywords):
    return [(d, k, f"{k} {d} İzmir") for d in districts for k in keywords]


def collect(client, queries, records):
    """Bütün sorğuları işlədir, nəticələri place id üzrə `records`-a yığır (təkrarlar birləşir)."""
    for i, (district, keyword, text_query) in enumerate(queries, 1):
        places = search_area(client, text_query)
        new = 0
        for place in places:
            place_id = place.get("id")
            if not place_id:
                continue
            rec = records.get(place_id)
            if rec is None:
                rec = records[place_id] = {"place": place, "keywords": {}, "query_district": district}
                new += 1
            rec["keywords"][keyword] = None  # dict: sıra qorunur, təkrar olmur
        log(f"[{i}/{len(queries)}] {text_query}: {len(places)} nəticə, {new} yeni")
    return records


def build_rows(records):
    labs, review = [], []
    skipped = {"closed": 0, "outside": 0}
    for rec in records.values():
        place = rec["place"]
        status = place.get("businessStatus")
        if status == "CLOSED_PERMANENTLY":
            skipped["closed"] += 1
            continue
        if not is_in_izmir(place):
            skipped["outside"] += 1
            continue
        location = place.get("location", {})
        row = {
            "name": place.get("displayName", {}).get("text", ""),
            "district": district_of(place, rec["query_district"]),
            "address": place.get("formattedAddress", ""),
            "phone": place.get("nationalPhoneNumber") or place.get("internationalPhoneNumber", ""),
            "website": place.get("websiteUri", ""),
            "rating": place.get("rating", ""),
            "reviews": place.get("userRatingCount", ""),
            "maps_url": place.get("googleMapsUri", ""),
            "category": place.get("primaryTypeDisplayName", {}).get("text") or place.get("primaryType", ""),
            "note": "Müvəqqəti bağlı" if status == "CLOSED_TEMPORARILY" else "",
            "lat": location.get("latitude", ""),
            "lng": location.get("longitude", ""),
            "keywords": ", ".join(rec["keywords"]),
        }
        (labs if is_likely_lab(row["name"]) else review).append(row)
    sort_key = lambda r: (normalize(r["district"]), normalize(r["name"]))  # noqa: E731
    labs.sort(key=sort_key)
    review.sort(key=sort_key)
    return labs, review, skipped


def write_csv(path, labs, review):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["Növ"] + [title for _, title in COLUMNS])
        for kind, rows in (("Lab", labs), ("Yoxlanmalı", review)):
            for row in rows:
                writer.writerow([kind] + [row[key] for key, _ in COLUMNS])


def write_xlsx(path, labs, review):
    wb = Workbook()
    for index, (title, rows) in enumerate((("Lablar", labs), ("Yoxlanmalı", review))):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        ws.append([t for _, t in COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row[key] for key, _ in COLUMNS])
        for col, (key, _) in enumerate(COLUMNS, 1):
            if key in LINK_FIELDS:
                for r in range(2, ws.max_row + 1):
                    cell = ws.cell(row=r, column=col)
                    if cell.value:
                        cell.hyperlink = cell.value
                        cell.style = "Hyperlink"
        widths = {"name": 40, "district": 14, "address": 60, "phone": 18, "website": 35,
                  "maps_url": 35, "category": 22, "keywords": 40}
        for col, (key, _) in enumerate(COLUMNS, 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = widths.get(key, 12)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def resolve_districts(names):
    if not names:
        return list(IZMIR_DISTRICTS)
    by_norm = {normalize(d): d for d in IZMIR_DISTRICTS}
    unknown = [n for n in names if normalize(n) not in by_norm]
    if unknown:
        raise SystemExit(
            f"Naməlum ilçə: {', '.join(unknown)}\nMövcud ilçələr: {', '.join(IZMIR_DISTRICTS)}"
        )
    return [by_norm[normalize(n)] for n in names]


def parse_args(argv):
    parser = argparse.ArgumentParser(
        description="İzmir diş protez laboratoriyalarını Google Places API ilə tapır."
    )
    parser.add_argument("--districts", nargs="+", metavar="İLÇƏ",
                        help="yalnız bu ilçələr (default: İzmirin bütün 30 ilçəsi)")
    parser.add_argument("--keywords", nargs="+", metavar="SÖZ",
                        help="açar sözlər (default: " + "; ".join(KEYWORDS) + ")")
    parser.add_argument("--output", default="output/izmir_dis_protez_lab",
                        help="çıxış faylının adı, uzantısız (default: %(default)s)")
    parser.add_argument("--dry-run", action="store_true",
                        help="API-yə getmədən sorğuları və təxmini sorğu sayını göstər")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    districts = resolve_districts(args.districts)
    queries = build_queries(districts, args.keywords or KEYWORDS)

    if args.dry_run:
        for _, _, text_query in queries:
            print(text_query)
        print(f"\n{len(queries)} sorğu, təxminən {len(queries)}–{len(queries) * MAX_PAGES} API çağırışı "
              "(sıx ilçələrin bölünməsi əlavə çağırış edə bilər)")
        return 0

    api_key = os.environ.get("GOOGLE_MAPS_API_KEY")
    if not api_key:
        log("GOOGLE_MAPS_API_KEY tapılmadı. Açarı environment variable kimi əlavə edin (README-yə baxın).")
        return 1

    client = PlacesClient(api_key)
    records = {}
    exit_code = 0
    try:
        collect(client, queries, records)
    except PlacesError as exc:
        log(f"\n{exc}")
        if not records:
            return 1
        log("Axtarış yarımçıq qaldı, indiyə qədər tapılanlar yazılır.")
        exit_code = 1

    labs, review, skipped = build_rows(records)
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    xlsx_path, csv_path = out.with_suffix(".xlsx"), out.with_suffix(".csv")
    write_xlsx(xlsx_path, labs, review)
    write_csv(csv_path, labs, review)

    counts = {}
    for kind, rows in (("lab", labs), ("review", review)):
        for row in rows:
            counts.setdefault(row["district"], {"lab": 0, "review": 0})[kind] += 1
    print("\nİlçə üzrə nəticə (lab / yoxlanmalı):")
    for district in sorted(counts, key=normalize):
        print(f"  {district:<14} {counts[district]['lab']:>4} / {counts[district]['review']}")
    print(f"\nCəmi: {len(labs)} lab, {len(review)} yoxlanmalı. "
          f"Atılan: {skipped['closed']} bağlı, {skipped['outside']} İzmirdən kənar.")
    print(f"API çağırışı: {client.request_count}")
    print(f"Fayllar: {xlsx_path}, {csv_path}")
    return exit_code


if __name__ == "__main__":
    sys.exit(main())
